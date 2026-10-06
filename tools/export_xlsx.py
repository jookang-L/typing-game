"""문제 은행(data/questions.json)을 검수용 엑셀로 내보낸다: python tools/export_xlsx.py
시트: 안내 / 타이핑 문제 / 퀴즈_<태그> 9개. 검수·메모 열은 비워 둔다."""
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).parent.parent
TAGS = ["출력", "변수", "입력", "연산자", "자료형", "컨테이너", "조건문", "반복문", "함수"]
LEVEL = {"square": "쉬움(정사각형)", "pyramid": "보통(피라미드)", "circle": "어려움(원)"}
LEVEL_ORDER = {"square": 0, "pyramid": 1, "circle": 2}
HEAD = PatternFill("solid", fgColor="1F4E5F")
INPUT = PatternFill("solid", fgColor="FFF2CC")


def sheet(wb, title, headers, rows, widths, wrap_cols=(), input_cols=()):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for r in rows:
        ws.append(r)
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEAD
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("=") and not cell.value.startswith("=INDEX("):
                cell.data_type = "s"   # '=' 나 '==' 로 시작하는 보기가 수식으로 바뀌지 않게 문자열로 고정
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column in wrap_cols)
            if cell.column in input_cols:
                cell.fill = INPUT
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return ws


def main():
    qs = json.loads((ROOT / "data" / "questions.json").read_text(encoding="utf-8"))
    qs.sort(key=lambda q: (TAGS.index(q["tag"]), LEVEL_ORDER[q["difficulty"]], q["id"]))
    wb = Workbook()
    ws = wb.active
    ws.title = "안내"
    for line in [
        "파이썬 타이핑 땅따먹기 문제 검수표",
        "",
        "구성: [타이핑 문제] 시트 1개 + [퀴즈_태그] 시트 9개. 문제 1개 = 명령어 1개 + 퀴즈 1개이고 ID로 연결됩니다.",
        "노란색 칸(검수, 수정 메모)에 적어 주세요. 검수: O(통과) / X(수정 필요) / 삭제",
        "수정할 내용은 해당 칸을 직접 고치거나 '수정 메모'에 적어 주세요. ID는 바꾸지 마세요.",
        "정답번호는 아래 보기 1~4 중 정답의 번호입니다. 게임에서는 보기 순서가 매번 섞입니다.",
        "명령어는 한 줄, 영문/숫자/기호만 사용합니다. 아이템 이름(bonus() x2 sleep(5) blur() try/except private hint())과 같으면 안 됩니다.",
        "퀴즈는 같은 행의 명령어와 같은 개념을 묻는 문제여야 합니다.",
        "주의: 보기를 '=' 나 '==' 로 시작하게 고칠 때는 앞에 작은따옴표(')를 붙여 입력하세요 (예: '==). 안 그러면 엑셀이 수식으로 바꿉니다.",
        "엑셀을 고친 뒤 게임에 반영하려면: python tools/import_xlsx.py review/문제검수.xlsx",
    ]:
        ws.append([line])
    ws["A1"].font = Font(bold=True, size=14)
    ws.column_dimensions["A"].width = 130

    sheet(wb, "타이핑 문제",
          ["ID", "태그", "난이도", "하위분류", "명령어", "글자수", "검수", "수정 메모"],
          [[q["id"], q["tag"], LEVEL[q["difficulty"]], q.get("sub", ""), q["command"], len(q["command"]), "", ""]
           for q in qs],
          [20, 10, 16, 12, 48, 8, 10, 40], input_cols=(7, 8))

    for tag in TAGS:
        rows = []
        for n, q in enumerate((x for x in qs if x["tag"] == tag), 2):
            z = q["quiz"]
            rows.append([q["id"], LEVEL[q["difficulty"]], q.get("sub", ""), z["type"], q["command"], z["question"],
                         *z["choices"], z["answer"] + 1, f"=INDEX(G{n}:J{n},K{n})", z["explanation"], "", ""])
        sheet(wb, f"퀴즈_{tag}",
              ["ID", "난이도", "하위분류", "퀴즈 유형", "명령어(참고)", "질문", "보기1", "보기2", "보기3", "보기4",
               "정답번호", "정답 내용", "해설", "검수", "수정 메모"],
              rows, [20, 15, 11, 11, 30, 52, 22, 22, 22, 22, 9, 22, 46, 9, 30],
              wrap_cols=(5, 6, 7, 8, 9, 10, 12, 13, 15), input_cols=(14, 15))
    out = ROOT / "review" / "문제검수.xlsx"
    wb.save(out)
    print(f"{len(qs)}세트 → {out}")


if __name__ == "__main__":
    main()
